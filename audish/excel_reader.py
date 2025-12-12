"""
Helper script to read Excel files and return JSON data for UI preview.
Can be called from Electron main process.
"""

import json
import sys
from pathlib import Path
from openpyxl import load_workbook


def read_excel_preview(file_path: str, max_rows: int = 100):
    """
    Read Excel file and return preview data as JSON.
    
    Args:
        file_path: Path to Excel file
        max_rows: Maximum number of rows to return
        
    Returns:
        JSON string with 'headers' and 'data' arrays
    """
    try:
        wb = load_workbook(file_path, data_only=True)
        sheet = wb.active
        
        # Get headers from first row
        headers = []
        for cell in sheet[1]:
            headers.append(str(cell.value) if cell.value is not None else '')
        
        # Get data rows
        data = []
        for row_idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_idx > max_rows + 1:
                break
            row_data = [str(cell) if cell is not None else '' for cell in row]
            # Pad row if needed
            while len(row_data) < len(headers):
                row_data.append('')
            data.append(row_data[:len(headers)])
        
        return json.dumps({
            'success': True,
            'headers': headers,
            'data': data,
            'total_rows': sheet.max_row - 1,
        })
    except Exception as e:
        return json.dumps({
            'success': False,
            'error': str(e),
        })


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(json.dumps({'success': False, 'error': 'Missing file path'}))
        sys.exit(1)
    
    file_path = sys.argv[1]
    max_rows = int(sys.argv[2]) if len(sys.argv) > 2 else 100
    
    result = read_excel_preview(file_path, max_rows)
    print(result)








