"""
Configuration validation with helpful, non-technical error messages.

This module validates all configuration files and data before scheduling,
catching problems early with actionable guidance for admissions staff.
"""

import yaml
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Set
from collections import defaultdict
from datetime import datetime
from difflib import get_close_matches


class ValidationError(Exception):
    """
    Raised when configuration is invalid.
    
    Contains a user-friendly error message with guidance on how to fix it.
    """
    pass


class ValidationResult:
    """Container for validation results with errors and warnings."""
    
    def __init__(self):
        self.errors: List[str] = []
        self.warnings: List[str] = []
    
    def add_error(self, message: str):
        """Add a fatal error."""
        self.errors.append(message)
    
    def add_warning(self, message: str):
        """Add a non-fatal warning."""
        self.warnings.append(message)
    
    @property
    def is_valid(self) -> bool:
        """True if no fatal errors."""
        return len(self.errors) == 0
    
    def raise_if_invalid(self):
        """Raise ValidationError if there are any errors."""
        if self.errors:
            raise ValidationError("\n\n".join(self.errors))
    
    def __str__(self) -> str:
        lines = []
        if self.errors:
            lines.append("❌ ERRORS (must fix before scheduling):")
            for err in self.errors:
                lines.append(f"\n{err}")
        if self.warnings:
            if lines:
                lines.append("\n")
            lines.append("⚠️  WARNINGS (schedule will run, but review these):")
            for warn in self.warnings:
                lines.append(f"\n  • {warn}")
        return "\n".join(lines) if lines else "✓ All validations passed"


def validate_all(
    mapping_path: str,
    rules_path: str,
    applicant_path: str,
    faculty_path: str,
) -> ValidationResult:
    """
    Run all validations and return combined results.
    
    This is the main entry point for validation. Run this before scheduling.
    
    Args:
        mapping_path: Path to mapping.yaml
        rules_path: Path to rules.yaml
        applicant_path: Path to applicant Excel file
        faculty_path: Path to faculty Excel file
        
    Returns:
        ValidationResult with all errors and warnings
    """
    result = ValidationResult()
    
    # Phase 1: Validate file existence
    _validate_files_exist(result, mapping_path, rules_path, applicant_path, faculty_path)
    if not result.is_valid:
        return result
    
    # Phase 2: Validate YAML syntax
    mapping_config = _validate_yaml_syntax(result, mapping_path, "mapping")
    rules_config = _validate_yaml_syntax(result, rules_path, "rules")
    if not result.is_valid:
        return result
    
    # Phase 3: Validate mapping structure
    _validate_mapping_structure(result, mapping_config, mapping_path)
    if not result.is_valid:
        return result
    
    # Phase 4: Validate Excel files match mapping
    app_headers = _validate_excel_file(result, applicant_path, mapping_config, 'applicants', mapping_path)
    fac_headers = _validate_excel_file(result, faculty_path, mapping_config, 'faculty', mapping_path)
    if not result.is_valid:
        return result
    
    # Phase 5: Validate rules structure
    _validate_rules_structure(result, rules_config, rules_path)
    if not result.is_valid:
        return result
    
    # Phase 6: Cross-validate (rules vs applicant data)
    if app_headers:
        _validate_rules_coverage(result, rules_config, applicant_path, mapping_config, rules_path)
    
    return result


def _validate_files_exist(
    result: ValidationResult,
    mapping_path: str,
    rules_path: str,
    applicant_path: str,
    faculty_path: str,
):
    """Check that all required files exist."""
    files = [
        (mapping_path, "Mapping configuration", "mapping.yaml"),
        (rules_path, "Rules configuration", "rules.yaml"),
        (applicant_path, "Applicant data", "Excel file"),
        (faculty_path, "Faculty availability", "Excel file"),
    ]
    
    for path, description, file_type in files:
        if not Path(path).exists():
            result.add_error(
                f"❌ File Not Found: {description}\n\n"
                f"Could not find: {path}\n\n"
                f"Please check that the {file_type} path is correct."
            )


def _validate_yaml_syntax(
    result: ValidationResult,
    path: str,
    config_name: str,
) -> Optional[Dict]:
    """Validate YAML file syntax and return parsed content."""
    try:
        with open(path) as f:
            content = yaml.safe_load(f)
            if content is None:
                result.add_error(
                    f"❌ Empty Configuration: {path}\n\n"
                    f"The {config_name} file is empty.\n\n"
                    f"Please add the required configuration."
                )
                return None
            return content
    except yaml.YAMLError as e:
        # Extract line number if available
        line_info = ""
        if hasattr(e, 'problem_mark') and e.problem_mark:
            line_info = f" (line {e.problem_mark.line + 1})"
        
        result.add_error(
            f"❌ Configuration Syntax Error: {path}{line_info}\n\n"
            f"The {config_name} file has invalid formatting.\n\n"
            f"Common causes:\n"
            f"  • Incorrect indentation (use spaces, not tabs)\n"
            f"  • Missing colon after a key\n"
            f"  • Unquoted special characters\n\n"
            f"Technical details: {str(e)[:200]}"
        )
        return None


def _validate_mapping_structure(
    result: ValidationResult,
    config: Dict,
    path: str,
):
    """Validate mapping.yaml has required sections and fields."""
    
    # Check for required sections
    if 'applicants' not in config:
        result.add_error(
            f"❌ Missing Section in {Path(path).name}\n\n"
            f"The 'applicants' section is required.\n\n"
            f"Add this to your mapping.yaml:\n\n"
            f"applicants:\n"
            f"  sheet: \"Export\"\n"
            f"  columns:\n"
            f"    id: \"Application Slate ID\"\n"
            f"    degree: \"Degree Level\"\n"
            f"    major: \"Major (Application)\""
        )
        return
    
    if 'faculty' not in config:
        result.add_error(
            f"❌ Missing Section in {Path(path).name}\n\n"
            f"The 'faculty' section is required.\n\n"
            f"Add this to your mapping.yaml:\n\n"
            f"faculty:\n"
            f"  sheet: \"Sheet1\"\n"
            f"  columns:\n"
            f"    faculty_name: \"Faculty\""
        )
        return
    
    # Check applicants section
    app_config = config['applicants']
    if not app_config.get('sheet'):
        result.add_error(
            f"❌ Missing Sheet Name in {Path(path).name}\n\n"
            f"The applicants section needs a 'sheet' name.\n\n"
            f"Fix: Add the sheet name:\n\n"
            f"applicants:\n"
            f"  sheet: \"Export\"  # ← Add this line"
        )
    
    # Check required applicant columns
    app_columns = app_config.get('columns', {})
    required_app_cols = {
        'id': 'Unique identifier for each applicant',
        'degree': 'Degree level (BM, MM, GD, AD, DMA)',
        'major': 'Major or discipline (Violin, Piano, etc.)',
    }
    
    missing_cols = [col for col in required_app_cols if col not in app_columns]
    if missing_cols:
        examples = "\n".join(f"    {col}: \"Column Name\"  # {desc}" 
                            for col, desc in required_app_cols.items() 
                            if col in missing_cols)
        result.add_error(
            f"❌ Missing Required Columns in {Path(path).name}\n\n"
            f"The applicants section is missing: {', '.join(missing_cols)}\n\n"
            f"Add these to applicants.columns:\n\n{examples}"
        )
    
    # Check faculty section
    fac_config = config['faculty']
    if not fac_config.get('sheet'):
        result.add_error(
            f"❌ Missing Sheet Name in {Path(path).name}\n\n"
            f"The faculty section needs a 'sheet' name.\n\n"
            f"Fix: Add the sheet name:\n\n"
            f"faculty:\n"
            f"  sheet: \"Sheet1\"  # ← Add this line"
        )
    
    # Check required faculty columns
    fac_columns = fac_config.get('columns', {})
    if 'faculty_name' not in fac_columns:
        result.add_error(
            f"❌ Missing Required Column in {Path(path).name}\n\n"
            f"The faculty section needs 'faculty_name' column.\n\n"
            f"Add this to faculty.columns:\n\n"
            f"    faculty_name: \"Faculty\"  # Column with faculty names"
        )


def _validate_excel_file(
    result: ValidationResult,
    excel_path: str,
    mapping_config: Dict,
    section: str,  # 'applicants' or 'faculty'
    mapping_path: str,
) -> Optional[List[str]]:
    """
    Validate Excel file matches mapping configuration.
    Returns list of headers if successful.
    """
    from openpyxl import load_workbook
    
    config = mapping_config.get(section, {})
    expected_sheet = config.get('sheet', 'Sheet1')
    columns = config.get('columns', {})
    
    try:
        wb = load_workbook(excel_path, read_only=True, data_only=True)
        available_sheets = wb.sheetnames
        
        # Check sheet exists
        # If there's only one sheet, use it regardless of name (handles copy/paste scenarios)
        if expected_sheet not in available_sheets:
            if len(available_sheets) == 1:
                # Single sheet in workbook - use it automatically
                # This handles the common case where users copy/paste data into a new workbook
                sheet = wb[available_sheets[0]]
                result.add_warning(
                    f"Sheet '{expected_sheet}' not found in {Path(excel_path).name}, "
                    f"using '{available_sheets[0]}' (only sheet in workbook)."
                )
            else:
                # Multiple sheets but expected one not found - this is an error
                suggestions = get_close_matches(expected_sheet, available_sheets, n=1, cutoff=0.5)
                
                suggestion_text = ""
                if suggestions:
                    suggestion_text = f"\n\nDid you mean '{suggestions[0]}'?"
                
                result.add_error(
                    f"❌ Sheet Not Found: '{expected_sheet}'\n\n"
                    f"File: {Path(excel_path).name}\n"
                    f"Available sheets: {', '.join(available_sheets)}{suggestion_text}\n\n"
                    f"Fix: Update {Path(mapping_path).name}:\n\n"
                    f"{section}:\n"
                    f"  sheet: \"{available_sheets[0]}\"  # Use exact sheet name"
                )
                wb.close()
                return None
        else:
            # Get headers from the expected sheet
            sheet = wb[expected_sheet]
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        wb.close()
        
        if not first_row:
            result.add_error(
                f"❌ Empty Sheet: '{expected_sheet}'\n\n"
                f"File: {Path(excel_path).name}\n\n"
                f"The sheet has no data. Please check the file."
            )
            return None
        
        headers = [str(h) if h else "" for h in first_row]
        
        # Check mapped columns exist in Excel
        missing_in_excel = []
        for logical_name, physical_name in columns.items():
            if physical_name not in headers:
                # Find similar column names
                suggestions = get_close_matches(physical_name, headers, n=1, cutoff=0.6)
                missing_in_excel.append((logical_name, physical_name, suggestions))
        
        if missing_in_excel:
            lines = [f"❌ Column Mismatch in {Path(excel_path).name}\n"]
            lines.append(f"These columns from {Path(mapping_path).name} weren't found:\n")
            
            for logical, physical, suggestions in missing_in_excel[:5]:
                if suggestions:
                    lines.append(f"  • '{physical}' → Did you mean '{suggestions[0]}'?")
                else:
                    lines.append(f"  • '{physical}'")
            
            if len(missing_in_excel) > 5:
                lines.append(f"  • ...and {len(missing_in_excel) - 5} more")
            
            lines.append(f"\nFix: Update the column names in {Path(mapping_path).name}")
            lines.append(f"to match exactly what's in the Excel file.")
            
            result.add_error("\n".join(lines))
            return None
        
        return headers
        
    except Exception as e:
        result.add_error(
            f"❌ Cannot Read File: {Path(excel_path).name}\n\n"
            f"Error: {str(e)[:200]}\n\n"
            f"Make sure the file is a valid Excel file (.xlsx) and isn't open in Excel."
        )
        return None


def _validate_rules_structure(
    result: ValidationResult,
    config: Dict,
    path: str,
):
    """Validate rules.yaml has required sections."""
    
    # Check for calendar section
    if 'calendar' not in config:
        result.add_error(
            f"❌ Missing Section in {Path(path).name}\n\n"
            f"The 'calendar' section is required.\n\n"
            f"Add this to your rules.yaml:\n\n"
            f"calendar:\n"
            f"  days:\n"
            f"    - {{ date: 2025-03-01, start: \"09:00\", end: \"17:00\" }}\n"
            f"    - {{ date: 2025-03-02, start: \"09:00\", end: \"17:00\" }}"
        )
        return
    
    calendar = config['calendar']
    if 'days' not in calendar or not calendar['days']:
        result.add_error(
            f"❌ No Audition Days Configured in {Path(path).name}\n\n"
            f"The calendar.days list is empty or missing.\n\n"
            f"Add audition days:\n\n"
            f"calendar:\n"
            f"  days:\n"
            f"    - {{ date: 2025-03-01, start: \"09:00\", end: \"17:00\" }}"
        )
        return
    
    # Validate each calendar day
    for i, day in enumerate(calendar['days']):
        if not isinstance(day, dict):
            result.add_error(
                f"❌ Invalid Day Format in {Path(path).name}\n\n"
                f"Calendar day #{i+1} is not properly formatted.\n\n"
                f"Each day should look like:\n"
                f"  - {{ date: 2025-03-01, start: \"09:00\", end: \"17:00\" }}"
            )
            continue
        
        if 'date' not in day:
            result.add_error(
                f"❌ Missing Date in {Path(path).name}\n\n"
                f"Calendar day #{i+1} is missing the 'date' field.\n\n"
                f"Add date to the day:\n"
                f"  - {{ date: 2025-03-01, start: \"09:00\", end: \"17:00\" }}"
            )
        
        # Check for start/end times (unless global defaults exist)
        has_global_start = calendar.get('default_start_time')
        has_global_end = calendar.get('default_end_time')
        
        if not has_global_start and 'start' not in day:
            result.add_warning(
                f"Day {day.get('date', f'#{i+1}')} has no start time. "
                f"Set 'start' or use global 'default_start_time'."
            )
        
        if not has_global_end and 'end' not in day:
            result.add_warning(
                f"Day {day.get('date', f'#{i+1}')} has no end time. "
                f"Set 'end' or use global 'default_end_time'."
            )
    
    # Check for rules section
    if 'rules' not in config:
        result.add_warning(
            f"No 'rules' section in {Path(path).name}. "
            f"Scheduling will fail if applicants don't match any rules."
        )
    
    # Check degree_precedence
    if 'degree_precedence' not in config:
        result.add_warning(
            f"No 'degree_precedence' in {Path(path).name}. "
            f"Using default: BM → MM → GD → AD → DMA"
        )


def _normalize_degree(degree: Optional[str]) -> Optional[str]:
    """
    Normalize degree values to standard codes.
    Matches the logic in mapping.py.
    """
    if not degree:
        return None
    
    degree_upper = str(degree).strip().upper()
    
    # Handle common variations
    if degree_upper in ['BM', 'BACHELOR OF MUSIC']:
        return 'BM'
    elif degree_upper in ['MM', 'MASTER OF MUSIC']:
        return 'MM'
    elif degree_upper in ['GD', 'GRADUATE DIPLOMA']:
        return 'GD'
    elif degree_upper in ['AD', 'ARTIST DIPLOMA']:
        return 'AD'
    elif degree_upper in ['DMA', 'DOCTOR OF MUSICAL ARTS']:
        return 'DMA'
    elif degree_upper in ['BCJ', 'BACHELOR OF JAZZ']:
        return 'BCJ'
    
    # Return as-is if no match
    return degree_upper


def _validate_rules_coverage(
    result: ValidationResult,
    rules_config: Dict,
    applicant_path: str,
    mapping_config: Dict,
    rules_path: str,
):
    """
    Check that rules exist for all discipline/degree combinations in applicant data.
    """
    from openpyxl import load_workbook
    
    app_config = mapping_config.get('applicants', {})
    sheet_name = app_config.get('sheet', 'Sheet1')
    columns = app_config.get('columns', {})
    
    degree_col = columns.get('degree')
    major_col = columns.get('major') or columns.get('department')
    
    if not degree_col or not major_col:
        return  # Can't validate without these columns
    
    try:
        wb = load_workbook(applicant_path, read_only=True, data_only=True)
        
        # Handle sheet name - use first sheet if only one available
        if sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
        elif len(wb.sheetnames) == 1:
            sheet = wb.active
        else:
            wb.close()
            return  # Sheet validation would have caught this already
        
        # Get header row to find column indices
        headers = [str(h) if h else "" for h in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        
        try:
            degree_idx = headers.index(degree_col)
            major_idx = headers.index(major_col)
        except ValueError:
            wb.close()
            return
        
        # Collect unique discipline/degree combinations (with normalized degrees)
        combos = defaultdict(int)  # (discipline, degree) -> count
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row and len(row) > max(degree_idx, major_idx):
                discipline = str(row[major_idx]) if row[major_idx] else None
                raw_degree = str(row[degree_idx]) if row[degree_idx] else None
                # Normalize degree to match rules.yaml format (BM, MM, etc.)
                degree = _normalize_degree(raw_degree)
                if discipline and degree:
                    combos[(discipline, degree)] += 1
        
        wb.close()
        
        # Check against rules
        rules = rules_config.get('rules', {})
        missing_rules = []
        
        for (discipline, degree), count in combos.items():
            if discipline not in rules:
                missing_rules.append((discipline, degree, count, 'discipline'))
            elif degree not in rules[discipline] and 'ALL' not in rules[discipline]:
                missing_rules.append((discipline, degree, count, 'degree'))
        
        if missing_rules:
            # Sort by count (most affected first)
            missing_rules.sort(key=lambda x: -x[2])
            
            lines = [f"❌ Missing Scheduling Rules in {Path(rules_path).name}\n"]
            lines.append("These applicant groups don't have scheduling rules:\n")
            
            # Get existing discipline names from rules for suggestions
            existing_disciplines = list(rules.keys())
            
            total_affected = 0
            for discipline, degree, count, missing_type in missing_rules[:7]:
                line = f"  • {discipline} - {degree}: {count} applicants"
                
                # Suggest similar discipline names
                if missing_type == 'discipline':
                    suggestions = get_close_matches(discipline, existing_disciplines, n=1, cutoff=0.5)
                    if suggestions:
                        line += f" (similar to: '{suggestions[0]}')"
                
                lines.append(line)
                total_affected += count
            
            if len(missing_rules) > 7:
                remaining = sum(x[2] for x in missing_rules[7:])
                lines.append(f"  • ...and {len(missing_rules) - 7} more ({remaining} applicants)")
                total_affected += remaining
            
            lines.append(f"\nTotal: {total_affected} applicants cannot be scheduled.")
            lines.append(f"\nFix: Add rules for each discipline in {Path(rules_path).name}:\n")
            
            # Show example fixes
            example_disciplines = set()
            for discipline, degree, _, _ in missing_rules[:3]:
                if discipline not in example_disciplines:
                    example_disciplines.add(discipline)
                    lines.append(f"  {discipline}:")
                    lines.append(f"    ALL: {{ cadence: {{ type: fixed_interval, minutes: 15 }} }}")
            
            result.add_error("\n".join(lines))
            
    except Exception as e:
        result.add_warning(f"Could not validate rules coverage: {str(e)[:100]}")


def validate_applicant_data(
    applicants: List[Dict],
    mapping_config: Dict,
) -> ValidationResult:
    """
    Validate loaded applicant data for common issues.
    
    Call this after loading and normalizing applicants.
    """
    result = ValidationResult()
    
    if not applicants:
        result.add_error(
            "❌ No Applicants Found\n\n"
            "The applicant file appears to be empty or has no valid data.\n\n"
            "Check that:\n"
            "  • The file has data rows (not just headers)\n"
            "  • The ID column contains values"
        )
        return result
    
    # Check for missing IDs
    missing_ids = [i+1 for i, a in enumerate(applicants) if not a.get('id')]
    if missing_ids:
        result.add_warning(
            f"{len(missing_ids)} applicants are missing IDs (rows: "
            f"{', '.join(str(r) for r in missing_ids[:5])}"
            f"{'...' if len(missing_ids) > 5 else ''}). These will be skipped."
        )
    
    # Check for missing degrees
    missing_degree = [a.get('id', f'row {i+1}') 
                      for i, a in enumerate(applicants) 
                      if a.get('id') and not a.get('degree')]
    if missing_degree:
        result.add_warning(
            f"{len(missing_degree)} applicants are missing degree information. "
            f"They may not match scheduling rules."
        )
    
    # Check for missing majors
    missing_major = [a.get('id', f'row {i+1}') 
                     for i, a in enumerate(applicants) 
                     if a.get('id') and not (a.get('major') or a.get('department'))]
    if missing_major:
        result.add_warning(
            f"{len(missing_major)} applicants are missing major/discipline. "
            f"They may not match scheduling rules."
        )
    
    # Check for duplicate IDs
    ids = [a.get('id') for a in applicants if a.get('id')]
    duplicates = [id for id in set(ids) if ids.count(id) > 1]
    if duplicates:
        result.add_error(
            f"❌ Duplicate Applicant IDs Found\n\n"
            f"These IDs appear more than once: {', '.join(str(d) for d in duplicates[:5])}"
            f"{'...' if len(duplicates) > 5 else ''}\n\n"
            f"Each applicant must have a unique ID. Check the source data."
        )
    
    return result


def validate_faculty_data(
    faculty_records: List[Dict],
    date_columns: List[str],
) -> ValidationResult:
    """
    Validate loaded faculty data.
    
    Call this after loading faculty records.
    """
    result = ValidationResult()
    
    if not faculty_records:
        result.add_error(
            "❌ No Faculty Records Found\n\n"
            "The faculty file appears to be empty.\n\n"
            "Check that the file has faculty data."
        )
        return result
    
    # Check for faculty without names
    unnamed = sum(1 for f in faculty_records if not f.get('faculty_name'))
    if unnamed:
        result.add_warning(
            f"{unnamed} faculty records are missing names and will be skipped."
        )
    
    # Check for date columns
    if not date_columns:
        result.add_error(
            "❌ No Date Columns Found\n\n"
            "Could not detect any date columns in the faculty file.\n\n"
            "Date columns should be formatted as dates (e.g., 2025-03-01)\n"
            "or recognizable date strings in the header row."
        )
    
    return result
